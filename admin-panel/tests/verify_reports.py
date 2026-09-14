"""Run full report/export regression flow against an isolated backup copy.

Usage: python -X utf8 tests/verify_reports.py
"""
import csv
import hashlib
import io
import os
import shutil
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from urllib.parse import parse_qs

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
SCRATCH = Path(tempfile.mkdtemp(prefix="aq_reports_flow_"))
BACKUP = ROOT.parent / ".backups/admin-panel-20260912-141835/aq_central.sqlite3"
shutil.copy2(BACKUP, SCRATCH / "test.sqlite3")
os.environ["AQ_DB"] = str(SCRATCH / "test.sqlite3")

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from app.main import app
from app.db import SessionLocal
from app.deps import Ctx, get_ctx
from app.models import User, Unit, Armory, ExportLog
from app.services import hisobot_svc as svc
from app.routers import jurnal

svc.EXPORT_DIR = SCRATCH / "exports"
jurnal.EXPORT_DIR = SCRATCH / "journal_exports"
with SessionLocal() as db:
    users = {u.username: u for u in db.scalars(select(User))}
    units = list(db.scalars(select(Unit)))
    region_armories = list(db.scalars(select(Armory.id).join(Unit, Unit.id == Armory.unit_id).where(Unit.region_id == 1)))
ctx = Ctx(user=users["admin"], lang="lat")
app.dependency_overrides[get_ctx] = lambda: ctx

with TestClient(app) as client:
    ids = []
    for slug in svc.REPORTS:
        for params in ({}, {"dan": "2026-09-12", "gacha": "2026-09-01", "hudud": "", "bolinma": ""},
                       {"hudud": "1", "bolinma": "1"}, {"page": "-2"}, {"page": "999999"}, {"hudud": "999999"}):
            r = client.get("/hisobotlar/" + slug, params=params)
            assert r.status_code == 200, (slug, params, r.status_code, r.text[:100])
            assert 1 <= r.context["page"] <= r.context["pages"]
            if params.get("hudud") == "999999":
                assert not r.context["rows"]
            if params.get("hudud") == "1":
                assert all(row.get("qurolxona_href") in ("", "/qurolxona/1") for row in r.context["rows"])
            if svc.REPORTS[slug]["dated"]:
                for choice in r.context["quick_dates"]:
                    q = parse_qs(choice["qs"])
                    start, end = datetime.fromisoformat(q["dan"][0]), datetime.fromisoformat(q["gacha"][0])
                    assert (end - start).days == choice["days"] - 1
        for fmt in ("csv", "xlsx", "pdf"):
            client.headers["X-CSRF-Token"] = client.cookies["aq_csrf"]
            r = client.post("/hisobotlar/" + slug + "/eksport", data={"fmt": fmt, "maqsad": "Test: " + slug, "muddat": "7",
                            "dan": "2026-08-01", "gacha": "2026-09-12", "hudud": "1"}, follow_redirects=False)
            assert r.status_code == 303, (slug, fmt, r.status_code, r.text)
            detail = client.get(r.headers["location"])
            assert detail.status_code == 200
            log_id = detail.context["log"].id
            r = client.get(f"/hisobotlar/eksportlar/{log_id}/fayl")
            assert r.status_code == 200, (slug, fmt, r.status_code)
            with SessionLocal() as db:
                log = db.get(ExportLog, log_id)
                ev = svc.export_event(db, log)
                assert log.file_hash == hashlib.sha256(r.content).hexdigest()
                assert ev and ev.payload["file_hash"] == log.file_hash and ev.payload["purpose"] == "Test: " + slug
                assert ev.payload["armory_ids"] == region_armories
                n = ev.payload["rows"]
            if fmt == "csv":
                csv_rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
                assert len(csv_rows)-1 == n
            if fmt == "xlsx":
                wb = load_workbook(io.BytesIO(r.content))
                assert wb["Hisobot"].max_row-6 == n
                assert wb["Eksport"]["B1"].value == log.number
                exported = [[svc.cell_text(value, col["kind"]) for value, col in zip(values, svc.REPORTS[slug]["columns"])]
                            for values in wb["Hisobot"].iter_rows(min_row=7, values_only=True)]
                # Administrator reports include the CSV export event created immediately
                # before XLSX. Other reports match cell-for-cell, allowing the CSV
                # protective apostrophe on text such as negative duration labels.
                if slug != "admin-harakatlari":
                    assert len(exported) == len(csv_rows) - 1
                    for row_index, (xlsx_row, csv_row) in enumerate(zip(exported, csv_rows[1:])):
                        assert len(xlsx_row) == len(csv_row)
                        for xlsx_text, csv_text in zip(xlsx_row, csv_row):
                            if csv_text == "'" + xlsx_text:
                                prefix = xlsx_text[:len(xlsx_text) - len(xlsx_text.lstrip())]
                                assert xlsx_text.lstrip().startswith(("=", "+", "-", "@")) or any(char in prefix for char in "\t\r\n")
                            else:
                                assert xlsx_text == csv_text, (slug, row_index, xlsx_text, csv_text)
            if fmt == "pdf":
                assert r.content.startswith(b"%PDF-") and b"%%EOF" in r.content
            ids.append(log_id)
        print(slug, "6 preview cases; CSV/XLSX/PDF created, audited, hashed and downloaded")
    assert client.get("/hisobotlar/texnik-holat?kesim=qurilma").status_code == 200
    assert client.get("/hisobotlar/eksportlar?dan=2026-09-12&gacha=2026-09-12&fmt=csv").status_code == 200
    for data, code in (({"fmt": "xml", "maqsad": "Test"}, "format"), ({"fmt": "csv", "maqsad": "   "}, "maqsad")):
        r = client.post("/hisobotlar/insidentlar/eksport", data=data, follow_redirects=False)
        assert r.status_code == 303 and "xato="+code in r.headers["location"]
    for unit in units:
        assert client.get(f"/hudud/{unit.region_id}/bolinma/{unit.id}").status_code == 200
    print("All", len(units), "unit detail pages OK; invalid export submission rejected")
    ctx.user = users["hudud_tk"]
    for slug in svc.REPORTS:
        r = client.get("/hisobotlar/"+slug+"?hudud=2&bolinma=1")
        assert r.status_code == 200
        assert r.context["f"].hudud == 1 and r.context["f"].bolinma == 1 and r.context["f"].armory_ids == [1]
    assert client.get(f"/hisobotlar/eksportlar/{ids[0]}/fayl").status_code == 200
    ctx.user = users["masul"]
    r = client.get("/hisobotlar/texnik-holat?hudud=2&bolinma=3")
    assert r.status_code == 200 and r.context["f"].armory_ids == [1]
    assert client.get(f"/hisobotlar/eksportlar/{ids[0]}").status_code == 403
    assert client.get(f"/hisobotlar/eksportlar/{ids[0]}/fayl").status_code == 403
    assert ids[0] not in [row["log"].id for row in client.get("/hisobotlar/eksportlar").context["rows"]]
    assert client.get("/hudud/1/bolinma/2").status_code == 403
    assert client.get("/hudud/2/bolinma/4").status_code == 403
    r = client.get("/hudud/1/bolinma/1")
    assert r.status_code == 200 and all(row["a"].id == 1 for row in r.context["armories"])
    r = client.post("/hisobotlar/texnik-holat/eksport", data={"fmt": "csv", "maqsad": "Own armory"}, follow_redirects=False)
    own = client.get(r.headers["location"]).context["log"].id
    assert client.get(f"/hisobotlar/eksportlar/{own}/fayl").status_code == 200
    r = client.post("/jurnal/eksport", data={"maqsad": "Journal integration check", "dan": "2026-09-01", "gacha": "2026-09-12"}, follow_redirects=False)
    assert r.status_code == 303, r.text[:200]
    journal_id = int(r.headers["location"].rsplit("/", 1)[1])
    native = client.get(f"/jurnal/eksport/{journal_id}/fayl")
    shared = client.get(f"/hisobotlar/eksportlar/{journal_id}/fayl")
    assert native.status_code == shared.status_code == 200 and native.content == shared.content
    detail = client.get(f"/hisobotlar/eksportlar/{journal_id}", follow_redirects=False)
    assert detail.status_code == 303 and detail.headers["location"] == f"/jurnal/eksport/{journal_id}"
    journal_rows = client.get("/hisobotlar/eksportlar?hisobot=jurnal").context["rows"]
    assert journal_rows and all(row["log"].report == "jurnal" for row in journal_rows)
    assert client.get("/hisobotlar/jurnal", follow_redirects=False).headers["location"] == "/jurnal"
    print("Journal exports shared path, download, detail redirect and list filter passed")
    ctx.user = User(username="foreign", full_name="Foreign", role="tekshiruvchi", scope_kind="qurolxona", scope_id=2, active=True)
    assert client.get(f"/hisobotlar/eksportlar/{own}/fayl").status_code == 403
    assert client.get(f"/hisobotlar/eksportlar/{journal_id}/fayl").status_code == 403
    ctx.user = users["admin"]
    with SessionLocal() as db:
        log = db.get(ExportLog, own)
        p = svc.file_path(log)
        data = p.read_bytes()
        p.write_bytes(data+b"tampered")
    assert client.get(f"/hisobotlar/eksportlar/{own}/fayl").status_code == 409
    p.write_bytes(data)
    with SessionLocal() as db:
        log = db.get(ExportLog, own)
        log.valid_until = datetime.now()-timedelta(seconds=1)
        db.commit()
    assert client.get(f"/hisobotlar/eksportlar/{own}/fayl").status_code == 410
    print("Scope isolation, own export, foreign export, hash mismatch and expiry checks OK")
print("PASS scratch:", SCRATCH)
