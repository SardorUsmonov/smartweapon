"""Regression checks for Unicode, full identifiers, literal text and export layout."""
import csv
import io
import os
import sys
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
OUT = Path(tempfile.mkdtemp(prefix="aq_export_renderers_"))
os.environ["AQ_DB"] = str(OUT / "unused.sqlite3")
from openpyxl import load_workbook
from pypdf import PdfReader
from app.services import hisobot_svc as svc
from app.i18n import to_cyrillic

cols = [dict(c, label=to_cyrillic(c["label"])) for c in svc.REPORTS["inventar-farqlari"]["columns"]]
row = {"vaqt": datetime(2026,9,12,10,15), "hudud": "Тошкент шаҳри", "bolinma": "Мирзо Улуғбек тумани ИИБ",
       "qurolxona": "Мирзо Улуғбек тумани ИИБ қуролхонаси", "yacheyka": "Y-001", "xodim": "Ғуломов Шерзод Ўткир ўғли",
       "tabel": "00123", "jihoz": "автомат АК-74", "seriya": "00123456789012345678", "rfid": "E200DEEF9C07A75114374509",
       "olindi": datetime(2026,9,12,8,20), "farq": "Текширувда аниқланган номутаносиблик: бириктирилган қурол RFID рақами қайтарилган қурол маълумотларига мос эмас. " * 3,
       "holat": ("chip-red", "Ечилмаган")}
meta = {"title": "Инвентар фарқлари", "period": "12.09.2026", "scope": "Тошкент шаҳри", "number": "VISUAL-CYR-001", "user": "Текширувчи",
        "ts": datetime.now(), "purpose": "Кирилл ёзувидаги узун матнлар ва идентификаторларни экспортда тўлиқ сақлашни текшириш. " * 2,
        "total": 1, "valid_until": datetime.now()+timedelta(days=7)}
pdf = svc.render_pdf(cols, [row], meta)
(OUT / "cyrillic-sample.pdf").write_bytes(pdf)
reader = PdfReader(io.BytesIO(pdf))
text = "\n".join(page.extract_text() for page in reader.pages)
assert len(reader.pages) == 1
for exact in ("Инвентар фарқлари", "00123", "00123456789012345678", "E200DEEF9C07A75114374509"):
    assert exact in text, ("PDF split/lost text", exact)
assert "\ufffd" not in text
xlsx = svc.render_xlsx(cols, [row], meta)
(OUT / "cyrillic-sample.xlsx").write_bytes(xlsx)
wb = load_workbook(io.BytesIO(xlsx))
ws = wb["Hisobot"]
assert ws["G7"].value == "00123" and ws["G7"].data_type == "s"
assert ws["I7"].value == "00123456789012345678" and ws["I7"].data_type == "s"
assert ws["A7"].value == row["vaqt"] and ws["A7"].number_format == "dd.mm.yyyy hh:mm"
assert ws["L7"].value == row["farq"] and ws["L7"].alignment.wrap_text
assert ws.row_dimensions[7].height > 60
assert wb["Eksport"]["B7"].value == meta["purpose"] and wb["Eksport"]["B7"].alignment.wrap_text
assert wb["Eksport"].row_dimensions[7].height > 30
assert ws.freeze_panes == "A7" and ws.auto_filter.ref == "A6:M7"
formula_case = svc.render_xlsx([svc.C("text", "Text"), svc.C("count", "Count", "n")], [{"text": '=HYPERLINK("https://example.com")', "count": 42}], meta)
with ZipFile(io.BytesIO(formula_case)) as archive:
    assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")
ws = load_workbook(io.BytesIO(formula_case))["Hisobot"]
assert ws["A7"].data_type == "s" and ws["A7"].value.startswith("=HYPERLINK")
assert ws["B7"].value == 42 and ws["B7"].data_type == "n"
csv_columns = [svc.C("text", "Text"), svc.C("count", "Count", "n"), svc.C("id", "ID", "mono")]
unsafe_text = ['=HYPERLINK("https://example.com")', "+SUM(1,2)", "-1+2", "@SUM(1,2)",
               "\ttext", "\rtext", "\ntext", "  =1+2", " \tplain", " \rplain", " \nplain", "\u00a0@SUM(1,2)"]
plain_text = ["Oddiy matn", "Текширув", "  ordinary text", "00123", "O'zbekiston, Toshkent", "Izoh\nkeyingi qator", ""]
csv_source = [{"text": value, "count": -42.5, "id": "00123456789012345678"} for value in unsafe_text + plain_text]
csv_source.append({"text": "numeric-looking user text", "count": "-42", "id": "00123"})
csv_data = svc.render_csv(csv_columns, csv_source, meta)
csv_rows = list(csv.reader(io.StringIO(csv_data.decode("utf-8-sig"))))
assert csv_rows[0] == ["Text", "Count", "ID"]
for index, original in enumerate(unsafe_text + plain_text, start=1):
    expected = "'" + original if index <= len(unsafe_text) else original
    assert csv_rows[index] == [expected, "-42.5", "00123456789012345678"], (original, csv_rows[index])
assert csv_rows[-1] == ["numeric-looking user text", "'-42", "00123"]
print("PASS: Unicode PDF and intact identifiers; typed dates/numbers; literal XLSX and formula-safe CSV text; wrapped descriptions/metadata.")
print("Artifacts:", OUT)
