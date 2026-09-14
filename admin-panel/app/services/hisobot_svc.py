"""Hisobotlar: 7 hisobotni qurish (oldindan ko'rish + eksport), CSV/XLSX/PDF, ExportLog raqamlash va fayl hash.

Har hisobot `build()` orqali bir xil shaklda qaytadi: ustunlar ro'yxati, qatorlar (dict), jami soni, xulosa.
Ustun turlari: l (matn, chapga), n (son, o'ngga), dt (sana-vaqt), chip ((css-klass, yorliq)), mono (seriya/hash).
"""
from __future__ import annotations

import csv
import hashlib
import io
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

from sqlalchemy import Integer, and_, cast, func, or_, select
from sqlalchemy.orm import Session

from .. import config
from ..models import (Alarm, Armory, Cabinet, Custody, Device, Event, ExportLog, Item, Officer, Region, SettingsAudit, Unit)
from .events import LEVEL_LABEL, record_event
from .kpi import armory_ids_for_scope
from .live import hub

PER_PAGE = 50
EXPORT_MAX_ROWS = 20000
EXPORT_DIR = config.DATA_DIR / "exports"
FORMATS = {
    "csv": ("CSV", "text/csv; charset=utf-8"),
    "xlsx": ("XLSX", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "pdf": ("PDF", "application/pdf"),
}
VALID_DAYS = (1, 7, 30, 90)
# administrator harakatlari hisobotiga kiradigan, insidentlarga kirmaydigan turlar
ADMIN_TYPES = {"eksport_qilindi", "sozlama_ozgardi", "huquq_berish_tasdiqlandi"}
TYPE_LABEL = {
    "auth_ok": "Kirish tasdiqlandi", "rad_etildi": "Rad etildi", "rad_etildi_takror": "Takroriy rad", "qulf_ochildi": "Qulf ochildi",
    "eshik_ochildi": "Eshik ochildi", "eshik_yopildi": "Eshik yopildi", "avtomatik_qulflandi": "Avtomatik qulflandi",
    "avtomat_olindi": "Avtomat olindi", "avtomat_qaytarildi": "Avtomat qaytarildi", "pm_olindi": "To'pponcha olindi",
    "pm_qaytarildi": "To'pponcha qaytarildi", "qaytarish_nomuvofiq": "Qaytarish nomuvofiq", "kechikish": "Kechikish",
    "eshik_ochiq_qoldi": "Eshik ochiq qoldi", "urilish": "Urilish", "buzish_urinishi": "Buzish urinishi", "lyuk_ochildi": "Lyuk ochildi",
    "quvvat_uzildi": "Quvvat uzildi", "quvvat_tiklandi": "Quvvat tiklandi", "batareya_past": "Batareya past", "aloqa_uzildi": "Aloqa uzildi",
    "aloqa_tiklandi": "Aloqa tiklandi", "sinxronlandi": "Sinxronlandi", "mexanik_kalit": "Mexanik kalit", "favqulodda_ochish": "Favqulodda ochish",
    "rejim_boshlandi": "Rejim boshlandi", "rejim_tugadi": "Rejim tugadi", "shkaf_bloklandi": "Yacheyka bloklandi",
    "shkaf_blokdan_chiqdi": "Blokdan chiqdi", "biriktirildi": "Biriktirildi", "ajratildi": "Ajratildi", "smena_ochildi": "Smena ochildi",
    "smena_yopildi": "Smena yopildi", "oz_tekshiruv_xato": "O'z-tekshiruv xatosi", "qulf_xatosi": "Qulf xatosi", "datchik_xatosi": "Datchik xatosi",
    "disk_toldi": "Disk to'ldi", "kamera_oqimi_uzildi": "Kamera oqimi uzildi", "kamera_oqimi_tiklandi": "Kamera oqimi tiklandi",
    "vaqt_sinxroni_buzildi": "Vaqt sinxroni buzildi", "server_yoq_rejimi": "Server yo'q rejimi", "huquq_berish_tasdiqlandi": "Huquq berish tasdiqlandi",
    "eksport_qilindi": "Eksport qilindi", "zaxira_nusxa": "Zaxira nusxa", "zaxira_nusxa_xato": "Zaxira nusxa xatosi", "texnik_xizmat": "Texnik xizmat",
    "inventarizatsiya": "Inventarizatsiya", "sozlama_ozgardi": "Sozlama o'zgardi",
}
DEVICE_KIND = {"server": "Server", "kommutator": "Kommutator", "ups": "UPS", "nas": "NAS", "kamera": "Kamera", "terminal": "Terminal", "kontroller": "Kontroller"}


def C(key: str, label: str, kind: str = "l", href: str | None = None) -> dict:
    return {"key": key, "label": label, "kind": kind, "href": href}


PLACE_COLS = [C("hudud", "Hudud", "l", "hudud_href"), C("bolinma", "Bo'linma"), C("qurolxona", "Qurolxona", "l", "qurolxona_href")]

REPORTS: dict[str, dict] = {
    "berish-qaytarish": {
        "title": "Berish-qaytarish jurnali", "icon": "swap", "dated": True, "days": 7,
        "desc": "Davr ichida yacheykalardan olingan va qaytarilgan qurollar: xodim, jihoz, muddat, mosligi.",
        "columns": [C("olindi", "Olindi", "dt"), C("qaytarildi", "Qaytarildi", "dt")] + PLACE_COLS + [
            C("yacheyka", "Yacheyka", "l", "yacheyka_href"), C("xodim", "Xodim", "l", "xodim_href"), C("tabel", "Tabel", "mono"),
            C("jihoz", "Jihoz", "l", "jihoz_href"), C("seriya", "Seriya", "mono"), C("muddat", "Muddat", "dt"), C("holat", "Holat", "chip")],
    },
    "qurollanganlik": {
        "title": "Joriy qurollanganlik", "icon": "gun", "dated": False, "days": 0,
        "desc": "Hozir xodimlar qo'lidagi qurollar (ochiq custody): kim, qaysi yacheykadan, qachon olgan, muddati.",
        "columns": PLACE_COLS + [C("yacheyka", "Yacheyka", "l", "yacheyka_href"), C("xodim", "Xodim", "l", "xodim_href"), C("tabel", "Tabel", "mono"),
                                 C("lavozim", "Lavozim"), C("jihoz", "Jihoz", "l", "jihoz_href"), C("seriya", "Seriya", "mono"),
                                 C("olindi", "Olindi", "dt"), C("muddat", "Muddat", "dt"), C("qoldi", "Qolgan vaqt", "n"), C("holat", "Holat", "chip")],
    },
    "insidentlar": {
        "title": "Insidentlar", "icon": "alert", "dated": True, "days": 7,
        "desc": "OGOHLANTIRISH, SIGNAL va XAVFSIZLIK darajasidagi hodisalar hamda ular bo'yicha signal holati.",
        "columns": [C("vaqt", "Vaqt", "dt"), C("daraja", "Daraja", "chip"), C("tur", "Tur", "l", "jurnal_href"), C("sarlavha", "Sarlavha")] + PLACE_COLS + [
            C("yacheyka", "Yacheyka", "l", "yacheyka_href"), C("xodim", "Xodim", "l", "xodim_href"), C("usul", "Usul", "mono"),
            C("natija", "Natija"), C("sabab", "Sabab"), C("signal", "Signal", "chip", "signal_href")],
    },
    "admin-harakatlari": {
        "title": "Administrator harakatlari", "icon": "shield", "dated": True, "days": 30,
        "desc": "Sozlamalar auditi (kim, qachon, eski → yangi) va XAVFSIZLIK darajasidagi hodisalar: huquq berish, eksport, bloklash, rejimlar.",
        "columns": [C("vaqt", "Vaqt", "dt"), C("manba", "Manba", "chip"), C("tur", "Tur", "l", "jurnal_href"), C("kim", "Kim"), C("tafsilot", "Tafsilot")] + PLACE_COLS + [
            C("tasdiq", "Tasdiqlagan"), C("natija", "Natija")],
    },
    "texnik-holat": {
        "title": "Texnik holat", "icon": "cpu", "dated": False, "days": 0,
        "desc": "Qurolxonalar va qurilmalar: aloqa, sinxron, UPS, kutilayotgan hodisalar, kontrollerlar, batareya, harorat.",
        "columns": PLACE_COLS + [C("holat", "Holat", "chip"), C("sinxron", "Oxirgi sinxron", "dt"), C("wan", "WAN"), C("ups", "UPS", "n"),
                                 C("kutilmoqda", "Kutilayotgan hodisalar", "n"), C("qurilma", "Qurilmalar onlayn", "n"), C("ogoh", "Ogohlantirish", "n"),
                                 C("qur_oflayn", "Qurilma oflayn", "n"), C("kontroller", "Kontrollerlar onlayn", "n"), C("batareya", "Batareya past", "n"),
                                 C("harorat", "O'rt. harorat", "n")],
        "columns_qurilma": PLACE_COLS + [C("qurilma_tur", "Qurilma turi"), C("qurilma", "Qurilma", "l", "qurilma_href"), C("holat", "Holat", "chip"),
                                         C("korsatkich", "Ko'rsatkichlar", "mono"), C("last_seen", "Oxirgi aloqa", "dt")],
    },
    "kechikishlar": {
        "title": "Kechikishlar", "icon": "clock", "dated": True, "days": 30,
        "desc": "Belgilangan muddatda qaytarilmagan qurollar: hali xodimda bo'lganlar va kech qaytarilganlar, kechikish davomiyligi.",
        "columns": [C("muddat", "Muddat", "dt"), C("olindi", "Olindi", "dt"), C("qaytarildi", "Qaytarildi", "dt")] + PLACE_COLS + [
            C("yacheyka", "Yacheyka", "l", "yacheyka_href"), C("xodim", "Xodim", "l", "xodim_href"), C("tabel", "Tabel", "mono"),
            C("jihoz", "Jihoz", "l", "jihoz_href"), C("seriya", "Seriya", "mono"), C("kechikish", "Kechikish", "n"), C("holat", "Holat", "chip")],
    },
    "inventar-farqlari": {
        "title": "Inventar farqlari", "icon": "clipboard", "dated": True, "days": 30,
        "desc": "Qaytarishda aniqlangan nomuvofiqliklar: RFID biriktirilgan qurolga mos kelmagan holatlar va ularning yechilishi.",
        "columns": [C("vaqt", "Aniqlandi", "dt")] + PLACE_COLS + [C("yacheyka", "Yacheyka", "l", "yacheyka_href"), C("xodim", "Xodim", "l", "xodim_href"),
                                                                 C("tabel", "Tabel", "mono"), C("jihoz", "Jihoz", "l", "jihoz_href"), C("seriya", "Seriya", "mono"),
                                                                 C("rfid", "RFID", "mono"), C("olindi", "Olindi", "dt"), C("farq", "Farq turi"), C("holat", "Holat", "chip")],
    },
}
REPORT_ORDER = list(REPORTS.keys())
EXPORT_REPORTS = {**REPORTS, "jurnal": {"title": "Hodisalar jurnali"}, "hudud_bolinmalar": {"title": "Hudud bo'linmalari"}}


# ---------------- Filtrlar ----------------
@dataclass
class Filt:
    dan: datetime | None = None
    gacha: datetime | None = None          # inklyuziv (kun oxiri sifatida ishlatiladi: < gacha + 1 kun)
    hudud: int | None = None
    bolinma: int | None = None
    kesim: str = ""
    armory_ids: list[int] | None = None    # None = cheklov yo'q
    scope_locked: str = ""                 # "hudud" | "bolinma" | "qurolxona" — foydalanuvchi doirasi

    @property
    def gacha_excl(self) -> datetime | None:
        return self.gacha + timedelta(days=1) if self.gacha else None

    def qs(self) -> dict:
        d = {}
        if self.dan:
            d["dan"] = self.dan.strftime("%Y-%m-%d")
        if self.gacha:
            d["gacha"] = self.gacha.strftime("%Y-%m-%d")
        if self.hudud:
            d["hudud"] = self.hudud
        if self.bolinma:
            d["bolinma"] = self.bolinma
        if self.kesim:
            d["kesim"] = self.kesim
        return d

    def period_label(self) -> str:
        if self.dan and self.gacha:
            return f"{self.dan:%d.%m.%Y} – {self.gacha:%d.%m.%Y}"
        return "joriy holat"


def _parse_date(s: str | None) -> datetime | None:
    if not s:
        return None
    try:
        return datetime.strptime(s.strip()[:10], "%Y-%m-%d")
    except ValueError:
        return None


def resolve_filters(db: Session, ctx, slug: str, dan: str = "", gacha: str = "", hudud: int | None = None,
                    bolinma: int | None = None, kesim: str = "", now: datetime | None = None) -> Filt:
    """So'rov parametrlari + foydalanuvchi vakolat doirasi -> Filt."""
    now = now or datetime.now()
    spec = REPORTS[slug]
    f = Filt(kesim=kesim or "")
    if spec["dated"]:
        f.dan = _parse_date(dan); f.gacha = _parse_date(gacha)
        if f.gacha is None:
            f.gacha = now.replace(hour=0, minute=0, second=0, microsecond=0)
        if f.dan is None:
            f.dan = f.gacha - timedelta(days=spec["days"] - 1)
        if f.dan > f.gacha:
            f.dan, f.gacha = f.gacha, f.dan
    # vakolat doirasi
    sk, sid = ctx.scope_kind, ctx.scope_id
    scope_ids = armory_ids_for_scope(db, sk, sid)
    if sk == "hudud" and sid:
        f.hudud = sid; f.scope_locked = "hudud"
        f.bolinma = bolinma or None
    elif sk == "bolinma" and sid:
        f.bolinma = sid; f.scope_locked = "bolinma"
        f.hudud = db.scalar(select(Unit.region_id).where(Unit.id == sid))
    elif sk == "qurolxona" and sid:
        f.scope_locked = "qurolxona"
        row = db.execute(select(Unit.id, Unit.region_id).join(Armory, Armory.unit_id == Unit.id).where(Armory.id == sid)).first()
        if row:
            f.bolinma, f.hudud = row[0], row[1]
    else:
        f.hudud = hudud or None
        f.bolinma = bolinma or None
    if f.bolinma and f.scope_locked not in ("bolinma", "qurolxona"):
        reg = db.scalar(select(Unit.region_id).where(Unit.id == f.bolinma))
        if reg is None:
            f.bolinma = None
        elif f.hudud and reg != f.hudud:
            f.bolinma = None   # bo'linma tanlangan hududga tegishli emas
        else:
            f.hudud = reg
    if f.hudud or f.bolinma or scope_ids is not None:
        q = select(Armory.id).join(Unit, Unit.id == Armory.unit_id)
        if f.hudud:
            q = q.where(Unit.region_id == f.hudud)
        if f.bolinma:
            q = q.where(Armory.unit_id == f.bolinma)
        if scope_ids is not None:
            q = q.where(Armory.id.in_(scope_ids))
        f.armory_ids = [r[0] for r in db.execute(q)]
    return f


def units_for(db: Session, region_id: int | None) -> list[Unit]:
    q = select(Unit).join(Region, Region.id == Unit.region_id).order_by(Region.order, Unit.name)
    if region_id:
        q = q.where(Unit.region_id == region_id)
    return list(db.execute(q).scalars())


def scope_label(db: Session, f: Filt) -> str:
    parts = []
    if f.hudud:
        r = db.get(Region, f.hudud)
        parts.append(r.name if r else str(f.hudud))
    if f.bolinma:
        u = db.get(Unit, f.bolinma)
        parts.append(u.name if u else str(f.bolinma))
    if f.scope_locked == "qurolxona" and f.armory_ids and len(f.armory_ids) == 1:
        a = db.get(Armory, f.armory_ids[0])
        parts.append(a.name if a else "")
    return " › ".join(p for p in parts if p) or "Respublika"


# ---------------- Yordamchilar ----------------
def _place(reg: Region, unit: Unit, arm: Armory, cab: Cabinet | None = None) -> dict:
    d = {"hudud": reg.short, "hudud_href": f"/hudud/{reg.id}", "bolinma": unit.name, "qurolxona": arm.name, "qurolxona_href": f"/qurolxona/{arm.id}"}
    if cab is not None:
        d.update(yacheyka=cab.label, yacheyka_href=f"/yacheyka/{cab.id}", _href=f"/yacheyka/{cab.id}")
    return d


def _officer(off: Officer | None) -> dict:
    if off is None:
        return {"xodim": "", "xodim_href": "", "tabel": "", "lavozim": ""}
    return {"xodim": off.full_name, "xodim_href": f"/xodim/{off.id}", "tabel": off.tabel, "lavozim": off.position}


def _item(it: Item | None) -> dict:
    if it is None:
        return {"jihoz": "", "jihoz_href": "", "seriya": "", "rfid": ""}
    return {"jihoz": f"{it.category} {it.model}".strip(), "jihoz_href": f"/jihoz/{it.id}", "seriya": it.serial, "rfid": it.rfid or ""}


def _custody_status(c: Custody, now: datetime) -> tuple[str, str]:
    if not c.match_ok:
        return ("chip-red", "Nomuvofiq")
    if c.returned_at is None:
        return ("chip-red", "Kechikkan") if c.due_at < now else ("chip-yellow", "Xodimda")
    if c.returned_at > c.due_at:
        return ("chip-yellow", "Kech qaytarildi")
    return ("chip-green", "Qaytarildi")


def _dur(delta: timedelta) -> str:
    s = int(delta.total_seconds())
    neg = s < 0
    s = abs(s)
    if s < 3600:
        out = "%d daq" % (s // 60)
    elif s < 86400:
        out = "%d soat %02d daq" % (s // 3600, (s % 3600) // 60)
    else:
        out = "%d kun %d soat" % (s // 86400, (s % 86400) // 3600)
    return ("-" if neg else "") + out


def _level_chip(level: str) -> tuple[str, str]:
    return {"CRITICAL": "chip-red", "SECURITY": "chip-red", "WARNING": "chip-yellow"}.get(level, "chip-gray"), LEVEL_LABEL.get(level, level)


def _join_custody(q):
    return (q.join(Item, Item.id == Custody.item_id).join(Officer, Officer.id == Custody.officer_id)
            .join(Cabinet, Cabinet.id == Custody.cabinet_id).join(Armory, Armory.id == Cabinet.armory_id)
            .join(Unit, Unit.id == Armory.unit_id).join(Region, Region.id == Unit.region_id))


def _join_event(q):
    return (q.outerjoin(Region, Region.id == Event.region_id).outerjoin(Unit, Unit.id == Event.unit_id)
            .outerjoin(Armory, Armory.id == Event.armory_id).outerjoin(Cabinet, Cabinet.id == Event.cabinet_id)
            .outerjoin(Officer, Officer.id == Event.officer_id).outerjoin(Alarm, Alarm.event_id == Event.id))


def _page(rows: list, offset: int, limit: int | None) -> list:
    return rows if limit is None else rows[offset:offset + limit]


# ---------------- Hisobotlar ----------------
def _rep_custody(db: Session, f: Filt, offset: int, limit: int | None, now: datetime, mode: str):
    """berish-qaytarish | qurollanganlik | kechikishlar | inventar-farqlari — Custody asosida."""
    where = []
    if f.armory_ids is not None:
        where.append(Armory.id.in_(f.armory_ids))
    if mode == "berish-qaytarish":
        where += [Custody.taken_at >= f.dan, Custody.taken_at < f.gacha_excl]
        order = Custody.taken_at.desc()
    elif mode == "qurollanganlik":
        where.append(Custody.returned_at.is_(None))
        order = Custody.due_at.asc()
    elif mode == "kechikishlar":
        where += [Custody.due_at >= f.dan, Custody.due_at < f.gacha_excl,
                  or_(and_(Custody.returned_at.is_(None), Custody.due_at < now), and_(Custody.returned_at.is_not(None), Custody.returned_at > Custody.due_at))]
        order = Custody.due_at.desc()
    else:  # inventar-farqlari
        where += [Custody.match_ok.is_(False), Custody.returned_at.is_not(None), Custody.returned_at >= f.dan, Custody.returned_at < f.gacha_excl]
        order = Custody.returned_at.desc()
    cq = _join_custody(select(func.count(Custody.id)).select_from(Custody)).where(*where)
    total = db.scalar(cq) or 0
    q = _join_custody(select(Custody, Item, Officer, Cabinet, Armory, Unit, Region)).where(*where).order_by(order, Custody.id.desc())
    q = q.offset(offset).limit(limit if limit is not None else EXPORT_MAX_ROWS)
    rows = []
    alarm_map: dict[int, list] = {}
    if mode == "inventar-farqlari":
        aq = select(Alarm.cabinet_id, Alarm.opened_at, Alarm.resolved_at).where(Alarm.type == "nomuvofiqlik")
        for cid, opened, resolved in db.execute(aq):
            alarm_map.setdefault(cid, []).append((opened, resolved))
    for c, it, off, cab, arm, unit, reg in db.execute(q):
        r = _place(reg, unit, arm, cab); r.update(_officer(off)); r.update(_item(it))
        r.update(olindi=c.taken_at, qaytarildi=c.returned_at, muddat=c.due_at, holat=_custody_status(c, now))
        if mode == "qurollanganlik":
            r["qoldi"] = _dur(c.due_at - now)
        if mode == "kechikishlar":
            r["kechikish"] = _dur((c.returned_at or now) - c.due_at)
            r["holat"] = ("chip-red", "Hali qaytarilmagan") if c.returned_at is None else ("chip-yellow", "Kech qaytarildi")
        if mode == "inventar-farqlari":
            r["vaqt"] = c.returned_at
            r["farq"] = "RFID biriktirilgan qurolga mos emas"
            resolved = None
            for opened, res in alarm_map.get(cab.id, []):
                if c.returned_at and abs((opened - c.returned_at).total_seconds()) <= 120:
                    resolved = res
            r["holat"] = ("chip-green", "Yechilgan") if resolved else ("chip-red", "Yechilmagan")
        rows.append(r)
    # xulosa
    summary = []
    if mode == "berish-qaytarish":
        open_n = db.scalar(cq.where(Custody.returned_at.is_(None))) or 0
        bad = db.scalar(cq.where(Custody.match_ok.is_(False))) or 0
        summary = [("Operatsiyalar", total, ""), ("Hali qaytarilmagan", open_n, "yellow" if open_n else ""), ("Nomuvofiq", bad, "red" if bad else "")]
    elif mode == "qurollanganlik":
        late = db.scalar(cq.where(Custody.due_at < now)) or 0
        summary = [("Xodimlarda", total, ""), ("Muddati o'tgan", late, "red" if late else "")]
    elif mode == "kechikishlar":
        still = db.scalar(cq.where(Custody.returned_at.is_(None))) or 0
        summary = [("Kechikishlar", total, "yellow" if total else ""), ("Hali qaytarilmagan", still, "red" if still else "")]
    else:
        unresolved = sum(1 for r in rows if r["holat"][1] == "Yechilmagan") if limit is None or total <= len(rows) else None
        summary = [("Farqlar", total, "red" if total else "")]
        if unresolved is not None:
            summary.append(("Yechilmagan", unresolved, "red" if unresolved else ""))
    return rows, total, summary


def _rep_incidents(db: Session, f: Filt, offset: int, limit: int | None, now: datetime):
    where = [Event.ts_server >= f.dan, Event.ts_server < f.gacha_excl, Event.level.in_(["WARNING", "CRITICAL", "SECURITY"]), Event.type.not_in(ADMIN_TYPES)]
    if f.armory_ids is not None:
        where.append(Event.armory_id.in_(f.armory_ids))
    total = db.scalar(select(func.count(Event.id)).where(*where)) or 0
    q = _join_event(select(Event, Region, Unit, Armory, Cabinet, Officer, Alarm)).where(*where).order_by(Event.ts_server.desc(), Event.id.desc())
    q = q.offset(offset).limit(limit if limit is not None else EXPORT_MAX_ROWS)
    rows = []
    for e, reg, unit, arm, cab, off, al in db.execute(q):
        r = {"vaqt": e.ts_server, "daraja": _level_chip(e.level), "tur": TYPE_LABEL.get(e.type, e.type), "jurnal_href": f"/jurnal/{e.id}",
             "sarlavha": e.title, "usul": e.method, "natija": e.result, "sabab": e.reason, "_href": f"/jurnal/{e.id}"}
        if reg is not None and unit is not None and arm is not None:
            r.update(_place(reg, unit, arm)); r["_href"] = f"/jurnal/{e.id}"
        else:
            r.update(hudud="Respublika", hudud_href="/", bolinma="", qurolxona="", qurolxona_href="")
        r.update(yacheyka=cab.label if cab else "", yacheyka_href=f"/yacheyka/{cab.id}" if cab else "")
        r.update(_officer(off))
        if al is None:
            r["signal"] = ("chip-gray", "—")
        elif al.resolved_at:
            r["signal"] = ("chip-green", "Yechilgan")
        elif al.acked_at:
            r["signal"] = ("chip-yellow", "Tasdiqlangan")
        else:
            r["signal"] = ("chip-red", "Faol")
        r["signal_href"] = f"/signal/{al.id}" if al else ""
        rows.append(r)
    by_level = dict(db.execute(select(Event.level, func.count(Event.id)).where(*where).group_by(Event.level)).all())
    summary = [("Insidentlar", total, ""), ("SIGNAL", by_level.get("CRITICAL", 0), "red" if by_level.get("CRITICAL") else ""),
               ("XAVFSIZLIK", by_level.get("SECURITY", 0), "red" if by_level.get("SECURITY") else ""),
               ("OGOHLANTIRISH", by_level.get("WARNING", 0), "yellow" if by_level.get("WARNING") else "")]
    return rows, total, summary


def _rep_admin(db: Session, f: Filt, offset: int, limit: int | None, now: datetime):
    rows = []
    # 1) sozlamalar auditi — markaziy harakatlar, faqat respublika kesimida
    n_audit = 0
    if f.armory_ids is None:
        aq = select(SettingsAudit).where(SettingsAudit.ts >= f.dan, SettingsAudit.ts < f.gacha_excl).order_by(SettingsAudit.ts.desc()).limit(EXPORT_MAX_ROWS)
        for s in db.execute(aq).scalars():
            n_audit += 1
            rows.append({"vaqt": s.ts, "manba": ("chip-yellow", "Sozlama"), "tur": "Sozlama o'zgardi", "jurnal_href": "/sozlamalar", "kim": s.username,
                         "tafsilot": f"{s.key}: {s.old or '—'} → {s.new or '—'}", "hudud": "Respublika", "hudud_href": "/", "bolinma": "", "qurolxona": "",
                         "qurolxona_href": "", "tasdiq": s.approved_by, "natija": "qo'llandi", "_href": "/sozlamalar"})
    # 2) XAVFSIZLIK darajasidagi hodisalar
    where = [Event.ts_server >= f.dan, Event.ts_server < f.gacha_excl, Event.level == "SECURITY"]
    if f.armory_ids is not None:
        where.append(Event.armory_id.in_(f.armory_ids))
    eq = _join_event(select(Event, Region, Unit, Armory, Cabinet, Officer, Alarm)).where(*where).order_by(Event.ts_server.desc()).limit(EXPORT_MAX_ROWS)
    n_ev = 0
    for e, reg, unit, arm, cab, off, _al in db.execute(eq):
        n_ev += 1
        who = e.approver1 or (off.full_name if off else "") or (e.payload or {}).get("user", "") or "tizim"
        r = {"vaqt": e.ts_server, "manba": ("chip-red", "Hodisa"), "tur": TYPE_LABEL.get(e.type, e.type), "jurnal_href": f"/jurnal/{e.id}", "kim": who,
             "tafsilot": (e.title + (" · " + e.detail if e.detail else "")) + ((" · " + cab.label) if cab else ""),
             "tasdiq": e.approver2, "natija": e.result, "_href": f"/jurnal/{e.id}"}
        if reg is not None and unit is not None and arm is not None:
            r.update(_place(reg, unit, arm))
        else:
            r.update(hudud="Respublika", hudud_href="/", bolinma="", qurolxona="", qurolxona_href="")
        rows.append(r)
    rows.sort(key=lambda r: r["vaqt"], reverse=True)
    total = len(rows)
    summary = [("Harakatlar", total, ""), ("Sozlamalar", n_audit, ""), ("Xavfsizlik hodisalari", n_ev, "red" if n_ev else "")]
    return _page(rows, offset, limit), total, summary


def _metrics_text(m: dict | None) -> str:
    if not m:
        return ""
    out = []
    for k, v in m.items():
        if isinstance(v, bool):
            out.append(f"{k}: {'ha' if v else 'yo‘q'}")
        elif isinstance(v, (int, float)) and k in ("cpu", "ram", "disk", "batareya", "yuklama"):
            out.append(f"{k} {v} %")
        else:
            out.append(f"{k} {v}")
    return " · ".join(out)


def _rep_tech(db: Session, f: Filt, offset: int, limit: int | None, now: datetime):
    if f.kesim == "qurilma":
        where = []
        if f.armory_ids is not None:
            where.append(Device.armory_id.in_(f.armory_ids))
        q = (select(Device, Armory, Unit, Region).join(Armory, Armory.id == Device.armory_id).join(Unit, Unit.id == Armory.unit_id)
             .join(Region, Region.id == Unit.region_id).where(*where))
        rows = []
        rank = {"oflayn": 0, "ogohlantirish": 1, "onlayn": 2}
        for d, arm, unit, reg in db.execute(q):
            r = _place(reg, unit, arm)
            chip = {"oflayn": ("chip-red", "Oflayn"), "ogohlantirish": ("chip-yellow", "Ogohlantirish")}.get(d.status, ("chip-green", "Onlayn"))
            r.update(qurilma_tur=DEVICE_KIND.get(d.kind, d.kind), qurilma=d.name, qurilma_href=f"/qurilma/{d.id}", holat=chip,
                     korsatkich=_metrics_text(d.metrics), last_seen=d.last_seen, _href=f"/qurilma/{d.id}", _rank=(rank.get(d.status, 3), reg.order, unit.name, d.kind))
            rows.append(r)
        rows.sort(key=lambda r: r["_rank"])
        total = len(rows)
        off_n = sum(1 for r in rows if r["holat"][1] == "Oflayn"); warn_n = sum(1 for r in rows if r["holat"][1] == "Ogohlantirish")
        summary = [("Qurilmalar", total, ""), ("Oflayn", off_n, "red" if off_n else ""), ("Ogohlantirish", warn_n, "yellow" if warn_n else "")]
        return _page(rows, offset, limit), total, summary
    # qurolxona kesimi
    where = []
    if f.armory_ids is not None:
        where.append(Armory.id.in_(f.armory_ids))
    dev = {}
    for aid, st, cnt in db.execute(select(Device.armory_id, Device.status, func.count(Device.id)).group_by(Device.armory_id, Device.status)):
        dev.setdefault(aid, {})[st] = cnt
    cabs = {}
    cq = select(Cabinet.armory_id, func.count(Cabinet.id), func.sum(cast(Cabinet.controller_online, Integer)),
                func.sum(cast(Cabinet.battery_pct < 20, Integer)), func.avg(Cabinet.temp_c)).group_by(Cabinet.armory_id)
    for aid, n, online, low, temp in db.execute(cq):
        cabs[aid] = (n or 0, int(online or 0), int(low or 0), float(temp or 0))
    q = select(Armory, Unit, Region).join(Unit, Unit.id == Armory.unit_id).join(Region, Region.id == Unit.region_id).where(*where)
    rows = []
    for arm, unit, reg in db.execute(q):
        d = dev.get(arm.id, {}); n_dev = sum(d.values()); n_on = d.get("onlayn", 0); n_warn = d.get("ogohlantirish", 0); n_off = d.get("oflayn", 0)
        n_cab, n_ctrl, n_low, temp = cabs.get(arm.id, (0, 0, 0, 0.0))
        if not arm.online:
            chip, rank = ("chip-red", "Oflayn"), 0
        elif n_off or n_warn or arm.ups_on_battery or arm.pending_events or (n_cab and n_ctrl < n_cab) or n_low:
            chip, rank = ("chip-yellow", "Ogohlantirish"), 1
        else:
            chip, rank = ("chip-green", "Me'yorda"), 2
        r = _place(reg, unit, arm)
        r.update(holat=chip, sinxron=arm.last_sync, wan="ha" if arm.wan_ok else "yo'q",
                 ups=f"{arm.ups_battery_pct} %" + (" · batareyada" if arm.ups_on_battery else ""), kutilmoqda=arm.pending_events,
                 qurilma=f"{n_on}/{n_dev}", ogoh=n_warn, qur_oflayn=n_off, kontroller=f"{n_ctrl}/{n_cab}", batareya=n_low,
                 harorat=f"{temp:.1f} °C" if n_cab else "", _href=f"/qurolxona/{arm.id}", _rank=(rank, reg.order, unit.name), _aid=arm.id)
        rows.append(r)
    rows.sort(key=lambda r: r["_rank"])
    total = len(rows)
    off_n = sum(1 for r in rows if r["holat"][1] == "Oflayn"); warn_n = sum(1 for r in rows if r["holat"][1] == "Ogohlantirish")
    n_dev_all = sum(sum(dev.get(r["_aid"], {}).values()) for r in rows)
    summary = [("Qurolxonalar", total, ""), ("Oflayn", off_n, "red" if off_n else ""), ("Ogohlantirish", warn_n, "yellow" if warn_n else ""), ("Qurilmalar", n_dev_all, "")]
    return _page(rows, offset, limit), total, summary


def columns_for(slug: str, f: Filt) -> list[dict]:
    spec = REPORTS[slug]
    if slug == "texnik-holat" and f.kesim == "qurilma":
        return spec["columns_qurilma"]
    return spec["columns"]


def build(db: Session, slug: str, f: Filt, page: int | None = 1, per_page: int = PER_PAGE, now: datetime | None = None):
    """Hisobotni quradi. page=None -> barcha qatorlar (eksport, EXPORT_MAX_ROWS gacha)."""
    now = now or datetime.now()
    offset = 0 if page is None else max(page - 1, 0) * per_page
    limit = None if page is None else per_page
    if slug in ("berish-qaytarish", "qurollanganlik", "kechikishlar", "inventar-farqlari"):
        rows, total, summary = _rep_custody(db, f, offset, limit, now, slug)
    elif slug == "insidentlar":
        rows, total, summary = _rep_incidents(db, f, offset, limit, now)
    elif slug == "admin-harakatlari":
        rows, total, summary = _rep_admin(db, f, offset, limit, now)
    elif slug == "texnik-holat":
        rows, total, summary = _rep_tech(db, f, offset, limit, now)
    else:
        raise KeyError(slug)
    return columns_for(slug, f), rows, total, summary


def overview(db: Session, f: Filt, now: datetime | None = None) -> dict[str, tuple[int, str, str]]:
    """Bosh sahifa kartochkalari uchun yengil ko'rsatkichlar: slug -> (son, izoh, ohang)."""
    now = now or datetime.now()
    week = now - timedelta(days=7)
    month = now - timedelta(days=30)

    def cust(*w):
        q = select(func.count(Custody.id)).join(Cabinet, Cabinet.id == Custody.cabinet_id).where(*w)
        if f.armory_ids is not None:
            q = q.where(Cabinet.armory_id.in_(f.armory_ids))
        return db.scalar(q) or 0

    def ev(*w):
        q = select(func.count(Event.id)).where(*w)
        if f.armory_ids is not None:
            q = q.where(Event.armory_id.in_(f.armory_ids))
        return db.scalar(q) or 0

    ops = cust(Custody.taken_at >= week)
    armed = cust(Custody.returned_at.is_(None))
    late = cust(Custody.returned_at.is_(None), Custody.due_at < now)
    inc = ev(Event.ts_server >= week, Event.level.in_(["WARNING", "CRITICAL", "SECURITY"]), Event.type.not_in(ADMIN_TYPES))
    adm = ev(Event.ts_server >= month, Event.level == "SECURITY")
    if f.armory_ids is None:
        adm += db.scalar(select(func.count(SettingsAudit.id)).where(SettingsAudit.ts >= month)) or 0
    aq = select(func.count(Armory.id)).where(Armory.online.is_(False))
    dq = select(func.count(Device.id)).where(Device.status != "onlayn")
    if f.armory_ids is not None:
        aq = aq.where(Armory.id.in_(f.armory_ids)); dq = dq.where(Device.armory_id.in_(f.armory_ids))
    off = db.scalar(aq) or 0; dev_bad = db.scalar(dq) or 0
    diff = cust(Custody.match_ok.is_(False), Custody.returned_at >= month)
    return {
        "berish-qaytarish": (ops, "operatsiya, 7 kun", ""),
        "qurollanganlik": (armed, "qurol hozir xodimlarda", ""),
        "insidentlar": (inc, "hodisa, 7 kun", "red" if inc else ""),
        "admin-harakatlari": (adm, "harakat, 30 kun", ""),
        "texnik-holat": (off, "qurolxona oflayn · %d qurilma muammoli" % dev_bad, "red" if off else ("yellow" if dev_bad else "")),
        "kechikishlar": (late, "hozir muddati o'tgan", "red" if late else ""),
        "inventar-farqlari": (diff, "nomuvofiqlik, 30 kun", "yellow" if diff else ""),
    }


# ---------------- Eksport ----------------
def cell_text(v, kind: str) -> str:
    if v is None:
        return ""
    if kind == "dt":
        return v.strftime("%d.%m.%Y %H:%M") if isinstance(v, datetime) else str(v)
    if kind == "chip":
        return v[1] if isinstance(v, (tuple, list)) else str(v)
    if isinstance(v, bool):
        return "ha" if v else "yo'q"
    return str(v)


def _table(cols: list[dict], rows: list[dict]) -> tuple[list[str], list[list[str]]]:
    header = [c["label"] for c in cols]
    body = [[cell_text(r.get(c["key"]), c["kind"]) for c in cols] for r in rows]
    return header, body


def render_csv(cols, rows, meta: dict) -> bytes:
    def safe_text(value: str) -> str:
        stripped = value.lstrip()
        leading = value[:len(value) - len(stripped)]
        if stripped.startswith(("=", "+", "-", "@")) or any(char in leading for char in "\t\r\n"):
            return "'" + value
        return value

    header = [safe_text(c["label"]) for c in cols]
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(header)
    for source in rows:
        values = []
        for col in cols:
            value = source.get(col["key"])
            text = cell_text(value, col["kind"])
            numeric = col["kind"] == "n" and isinstance(value, (int, float)) and not isinstance(value, bool)
            # CSV quoting does not prevent spreadsheet formula execution.
            values.append(text if numeric else safe_text(text))
        w.writerow(values)
    return buf.getvalue().encode("utf-8-sig")


def render_xlsx(cols, rows, meta: dict) -> bytes:
    import textwrap
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    header, body = _table(cols, rows)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hisobot"
    ws.append([meta["title"]]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Davr: {meta['period']} · Qamrov: {meta['scope']}"])
    ws.append([f"Eksport: {meta['number']} · {meta['user']} · {meta['ts']:%d.%m.%Y %H:%M} · Maqsad: {meta['purpose']}"])
    ws.append([f"Qatorlar: {meta['total']}" + (f" (faylda {len(body)} ta)" if meta["total"] != len(body) else "") + " · NAMUNAVIY MA'LUMOT (simulyator)"])
    ws.append([])
    ws.append(header)
    hrow = ws.max_row
    fill = PatternFill("solid", fgColor="0F1A2E")
    for cell in ws[hrow]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF"); cell.fill = fill; cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[hrow].height = 24
    for source in rows:
        row_number = ws.max_row + 1
        for i, col in enumerate(cols, start=1):
            value = source.get(col["key"])
            cell = ws.cell(row=row_number, column=i)
            if col["kind"] == "dt" and isinstance(value, datetime):
                cell.value = value
                cell.number_format = "dd.mm.yyyy hh:mm"
            elif col["kind"] == "n" and isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.value = value
            else:
                cell.value = cell_text(value, col["kind"])
                cell.data_type = "s"  # IDs retain leading zeros; user text cannot become a formula.
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", horizontal="right" if cell.data_type == "n" else "left", wrap_text=col["kind"] != "mono")
    ws.freeze_panes = ws.cell(row=hrow + 1, column=1)
    for i, h in enumerate(header, start=1):
        width = max([len(h)] + [len(str(r[i - 1])) for r in body]) if body else len(h)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 4, 12), 64) if cols[i - 1]["kind"] == "mono" else min(max(width + 2, 8), 48)
    for row_number, values in enumerate(body, start=hrow + 1):
        lines = max(1 if cols[i - 1]["kind"] == "mono" else sum(max(1, len(textwrap.wrap(part, width=max(1, int(ws.column_dimensions[get_column_letter(i)].width) - 3))))
                        for part in value.splitlines() or [""])
                    for i, value in enumerate(values, start=1))
        ws.row_dimensions[row_number].height = min(409, max(20, lines * 15 + 4))
    ws.auto_filter.ref = f"A{hrow}:{get_column_letter(len(header))}{max(ws.max_row, hrow)}"
    info = wb.create_sheet("Eksport")
    for k, v in (("Raqam", meta["number"]), ("Hisobot", meta["title"]), ("Davr", meta["period"]), ("Qamrov", meta["scope"]),
                 ("Kim", meta["user"]), ("Vaqt", f"{meta['ts']:%d.%m.%Y %H:%M:%S}"), ("Maqsad", meta["purpose"]),
                 ("Amal qilish muddati", f"{meta['valid_until']:%d.%m.%Y %H:%M}"), ("Manba", "Aqlli qurolxona · Nazorat markazi (prototip)")):
        info.append([k, v])
    info.column_dimensions["A"].width = 22; info.column_dimensions["B"].width = 60
    for row in info:
        for cell in row:
            cell.data_type = "s"
            cell.font = Font(name="Arial", size=10, bold=cell.column == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        lines = max(len(textwrap.wrap(str(row[0].value or ""), width=19)), len(textwrap.wrap(str(row[1].value or ""), width=57)), 1)
        info.row_dimensions[row[0].row].height = min(409, max(20, lines * 15 + 4))
    out = io.BytesIO(); wb.save(out)
    return out.getvalue()


_PDF_FONT: tuple[str, str] | None = None


def _pdf_fonts() -> tuple[str, str]:
    """Kirill va o'zbek lotin belgilari uchun TTF shrift (Arial/DejaVu); topilmasa Helvetica."""
    global _PDF_FONT
    if _PDF_FONT:
        return _PDF_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    cands = [("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
             ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
             ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf")]
    for reg, bold in cands:
        if Path(reg).exists():
            try:
                pdfmetrics.registerFont(TTFont("AQSans", reg))
                if Path(bold).exists():
                    pdfmetrics.registerFont(TTFont("AQSans-Bold", bold))
                    _PDF_FONT = ("AQSans", "AQSans-Bold")
                else:
                    _PDF_FONT = ("AQSans", "AQSans")
                return _PDF_FONT
            except Exception:  # noqa: BLE001
                continue
    _PDF_FONT = ("Helvetica", "Helvetica-Bold")
    return _PDF_FONT


def render_pdf(cols, rows, meta: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    font, bold = _pdf_fonts()
    header, body = _table(cols, rows)
    page_w, page_h = landscape(A4)
    margin = 24
    avail = page_w - 2 * margin
    st_title = ParagraphStyle("t", fontName=bold, fontSize=13, leading=16)
    st_meta = ParagraphStyle("m", fontName=font, fontSize=8.5, leading=11, textColor=colors.HexColor("#3a4a66"))
    st_cell = ParagraphStyle("c", fontName=font, fontSize=7.2, leading=8.6)
    st_head = ParagraphStyle("h", fontName=bold, fontSize=7.2, leading=8.6, textColor=colors.white)
    # Minimal width keeps ordinary words, dates and identifiers intact. Remaining
    # width goes to descriptive columns instead of shrinking short numeric IDs.
    minimums, desired = [], []
    for i, h in enumerate(header):
        text = [r[i] for r in body]
        kind = cols[i]["kind"]
        words = [word for value in text for word in value.split()]
        heading_min = max((stringWidth(word, bold, 7.2) for word in h.split()), default=0)
        value_min = max((stringWidth(word, font, 7.2) for word in words), default=0)
        if kind != "mono":
            value_min = min(value_min, 90)
        minimum = max(heading_min, value_min, 16) + 8
        full = max([stringWidth(h, bold, 7.2)] + [stringWidth(value, font, 7.2) for value in text])
        minimums.append(minimum)
        desired.append(max(minimum, min(full + 8, 155)))
    if sum(minimums) > avail:
        # Unusually long identifiers still wrap safely inside the existing page.
        widths = [avail * width / sum(minimums) for width in minimums]
    elif sum(desired) <= avail:
        widths = [avail * width / sum(desired) for width in desired]
    else:
        extra = avail - sum(minimums)
        capacity = sum(wanted - minimum for wanted, minimum in zip(desired, minimums))
        widths = [minimum + extra * (wanted - minimum) / capacity for wanted, minimum in zip(desired, minimums)]
    data = [[Paragraph(escape(h), st_head) for h in header]]
    for r in body:
        data.append([Paragraph(escape(v), st_cell) for v in r])
    tbl = Table(data, colWidths=widths, repeatRows=1)
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f1a2e")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b8c2d3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5fa")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
    ]))
    story = [Paragraph(escape(f"{meta['title']}"), st_title),
             Paragraph(escape(f"O'zbekiston Respublikasi IIV · Aqlli qurolxona · Nazorat markazi · Davr: {meta['period']} · Qamrov: {meta['scope']}"), st_meta),
             Paragraph(escape(f"Eksport {meta['number']} · {meta['user']} · {meta['ts']:%d.%m.%Y %H:%M} · Maqsad: {meta['purpose']} · Amal qilish muddati: {meta['valid_until']:%d.%m.%Y}"), st_meta),
             Paragraph(escape(f"Qatorlar: {meta['total']}" + (f" (hujjatda {len(body)} ta)" if meta['total'] != len(body) else "")), st_meta),
             Spacer(1, 6), tbl if body else Paragraph("Ma'lumot yo'q", st_meta)]

    def on_page(canvas, doc):
        canvas.saveState()
        canvas.setFont(font, 7.5); canvas.setFillColor(colors.HexColor("#5f7194"))
        canvas.drawString(margin, 14, f"{meta['number']} · NAMUNAVIY MA'LUMOT (simulyator) · hujjat hash raqami eksport jurnalida")
        canvas.drawRightString(page_w - margin, 14, f"Sahifa {doc.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=margin, rightMargin=margin, topMargin=22, bottomMargin=26,
                            title=meta["title"], author=meta["user"], subject=meta["number"])
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


RENDERERS = {"csv": render_csv, "xlsx": render_xlsx, "pdf": render_pdf}


def next_number(db: Session, now: datetime) -> str:
    """E-2026-000123: yil ichida ketma-ket raqam."""
    prefix = f"E-{now.year}-"
    n = (db.scalar(select(func.count(ExportLog.id)).where(ExportLog.number.like(prefix + "%"))) or 0) + 1
    while db.scalar(select(func.count(ExportLog.id)).where(ExportLog.number == f"{prefix}{n:06d}")):
        n += 1
    return f"{prefix}{n:06d}"


def file_path(log: ExportLog) -> Path:
    if log.report == "jurnal":
        from ..routers.jurnal import _export_path
        return _export_path(log)
    return EXPORT_DIR / f"{log.number}.{log.fmt}"


def file_status(log: ExportLog) -> tuple[str, str]:
    """Fayl mavjudligi va hash mosligi: (chip klass, yorliq)."""
    p = file_path(log)
    if not p.exists():
        return ("chip-red", "Fayl yo'q")
    h = hashlib.sha256(p.read_bytes()).hexdigest()
    return ("chip-green", "Hash mos") if h == log.file_hash else ("chip-red", "Hash mos emas")


def create_export(db: Session, ctx, slug: str, fmt: str, purpose: str, valid_days: int, f: Filt, now: datetime | None = None) -> ExportLog:
    """Faylni hosil qiladi, diskka yozadi, ExportLog (raqam, kim, maqsad, muddat, hash) va `eksport_qilindi` hodisasini yozadi."""
    now = (now or datetime.now()).replace(microsecond=0)
    if fmt not in RENDERERS:
        raise ValueError("format")
    if not purpose.strip():
        raise ValueError("maqsad")
    spec = REPORTS[slug]
    cols, rows, total, _summary = build(db, slug, f, page=None, now=now)
    username = ctx.user.username if ctx and ctx.user else "mehmon"
    number = next_number(db, now)
    valid_until = now + timedelta(days=int(valid_days) if int(valid_days) in VALID_DAYS else 7)
    meta = {"title": spec["title"], "period": f.period_label(), "scope": scope_label(db, f), "number": number, "user": username,
            "ts": now, "purpose": purpose.strip(), "total": total, "valid_until": valid_until}
    data = RENDERERS[fmt](cols, rows, meta)
    h = hashlib.sha256(data).hexdigest()
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    log = ExportLog(number=number, username=username, report=slug, fmt=fmt, purpose=purpose.strip()[:200], valid_until=valid_until, ts=now, file_hash=h)
    (EXPORT_DIR / f"{number}.{fmt}").write_bytes(data)
    db.add(log); db.flush()
    armory = db.get(Armory, f.armory_ids[0]) if f.armory_ids and len(f.armory_ids) == 1 else None
    ev = record_event(db, "eksport_qilindi", armory=armory, title="Hisobot eksport qilindi",
                      detail=f"{number} · {spec['title']} · {fmt.upper()} · {username}", approver1=username, simulated=False,
                      payload={"number": number, "report": slug, "fmt": fmt, "rows": min(total, len(rows)), "total": total, "purpose": purpose.strip()[:200],
                               "valid_until": valid_until.isoformat(), "file_hash": h, "size": len(data), "filters": f.qs(), "scope": meta["scope"],
                               "armory_ids": f.armory_ids})
    hub.broadcast_threadsafe({"type": "event", "event_type": ev.type, "level": ev.level, "cabinet_id": None, "armory_id": armory.id if armory else None,
                              "title": ev.title, "ts": ev.ts_server.isoformat()})
    return log


def export_event(db: Session, log: ExportLog) -> Event | None:
    """Eksportga tegishli hodisa; eski hudud eksportlari `raqam` kalitidan foydalanadi."""
    return db.execute(select(Event).where(Event.type == "eksport_qilindi", or_(
        Event.payload["number"].as_string() == log.number, Event.payload["raqam"].as_string() == log.number,
        Event.detail.like(log.number + " ·%"))).order_by(Event.id.desc())).scalars().first()


def can_view_export(db: Session, ctx, log: ExportLog) -> bool:
    """Eksport ichidagi butun qamrov joriy foydalanuvchi vakolatiga sig'ishi shart.

    Yangi eksport uchun jurnalga yozilgan yacheyka omborlari ro'yxati ishlatiladi.
    Eski yozuvning qamrovi aniqlanmasa, cheklangan foydalanuvchiga ko'rsatilmaydi.
    """
    if ctx.scope_kind == "respublika":
        return True
    allowed = set(armory_ids_for_scope(db, ctx.scope_kind, ctx.scope_id) or [])
    event = export_event(db, log)
    if event is None:
        return False
    payload = event.payload or {}
    if "armory_ids" in payload:
        exported = payload["armory_ids"]
        if not isinstance(exported, list) or any(type(aid) is not int for aid in exported):
            return False
        if not exported:
            return bool(ctx.user and log.username == ctx.user.username)
        return set(exported).issubset(allowed)
    filters = payload.get("filters") or {}
    if not isinstance(filters, dict):
        return False
    unit_id = filters.get("bolinma")
    region_id = filters.get("hudud") or payload.get("hudud_id")
    if unit_id or region_id:
        query = select(Armory.id).join(Unit, Unit.id == Armory.unit_id)
        if unit_id:
            query = query.where(Unit.id == unit_id)
        if region_id:
            query = query.where(Unit.region_id == region_id)
        exported = set(db.scalars(query))
        return bool(exported) and exported.issubset(allowed)
    return event.armory_id is not None and event.armory_id in allowed


def visible_export_ids(db: Session, ctx) -> list[int] | None:
    if ctx.scope_kind == "respublika":
        return None
    return [log.id for log in db.scalars(select(ExportLog)) if can_view_export(db, ctx, log)]
